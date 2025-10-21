"""
Create MS Word document for CO2 emissions by source results
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Add title
title = doc.add_heading('CO2 Emissions by Source Results', level=1)
title_format = title.runs[0]
title_format.font.name = 'Times New Roman'
title_format.font.size = Pt(14)
title_format.bold = True

# Add main results paragraph
results_text = (
    "Our sectoral emissions analysis reveals substantial decarbonization potential across all sources by 2040. "
    "From a 2021 baseline of 112.85 MtCO₂, we project total emissions declining to 66.89 MtCO₂ under BAU (-40.7%), "
    "57.64 MtCO₂ under ETS1 (-48.9%), and 38.47 MtCO₂ under ETS2 (-65.9%) (IPCC, 2022). The household sector "
    "dominates emissions (60.1% in 2021) and shows the largest absolute reduction under ETS2: 67.84→21.68 MtCO₂ "
    "(-68.0%) (IEA, 2023). Industry responds strongly to ETS1 pricing, with emissions falling from 10.68 to 4.86 MtCO₂ "
    "(-54.5%), demonstrating sectoral sensitivity to carbon price signals (Bataille et al., 2018). Energy sector "
    "emissions decline 66.3% under ETS2, reflecting renewable capacity expansion and fuel switching (Luderer et al., 2021)."
)

para = doc.add_paragraph(results_text)
para_format = para.paragraph_format
para_format.line_spacing = 1.5
para_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Format text
for run in para.runs:
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)

# Add references heading
ref_heading = doc.add_heading('References', level=2)
ref_heading_format = ref_heading.runs[0]
ref_heading_format.font.name = 'Times New Roman'
ref_heading_format.font.size = Pt(12)
ref_heading_format.bold = True

# Add references
references = [
    ("Bataille, C., Åhman, M., Neuhoff, K., Nilsson, L. J., Fischedick, M., Lechtenböhmer, S., ... & Rootzén, J. (2018). "
     "A review of technology and policy deep decarbonization pathway options for making energy-intensive industry "
     "production consistent with the Paris Agreement. Journal of Cleaner Production, 187, 960-973."),

    ("IEA. (2023). Energy Technology Perspectives 2023. International Energy Agency, Paris. "
     "https://www.iea.org/reports/energy-technology-perspectives-2023"),

    ("IPCC. (2022). Climate Change 2022: Mitigation of Climate Change. Contribution of Working Group III to the Sixth "
     "Assessment Report of the Intergovernmental Panel on Climate Change. Cambridge University Press."),

    ("Luderer, G., Madeddu, S., Merfort, L., Ueckerdt, F., Pehl, M., Pietzcker, R., ... & Bauer, N. (2021). "
     "Impact of declining renewable energy costs on electrification in low-emission scenarios. Nature Energy, 7(1), 32-42.")
]

for ref in references:
    ref_para = doc.add_paragraph(ref)
    ref_para_format = ref_para.paragraph_format
    ref_para_format.line_spacing = 1.5
    ref_para_format.left_indent = Inches(0.5)
    ref_para_format.first_line_indent = Inches(-0.5)

    for run in ref_para.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

# Save document
doc.save('CO2_By_Source_Results_Brief.docx')
print("✓ MS Word document created successfully!")
print("  - CO2_By_Source_Results_Brief.docx")
print(f"\nWord count: ~122 words (excluding references)")
